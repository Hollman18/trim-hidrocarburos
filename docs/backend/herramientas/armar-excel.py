"""Arma modelo-de-datos.xlsx a partir de tablas.json (la especificación del modelo de datos).

    pip install openpyxl pillow   (en un entorno virtual)
    python armar-excel.py docs/backend/sql/tablas.json docs/backend/modelo-de-datos.xlsx <diagramas>/er-*.png

Hojas: Léame · Tablas · Columnas · Relaciones · Filas de ejemplo · Diagrama (si se pasan imágenes).
Los conteos de la hoja Tablas son fórmulas (COUNTIF sobre Columnas y Relaciones), no números fijos.
Sale con error si la especificación se contradice (FK a una tabla o columna que no existe,
relación que nombra una columna que no está, tabla repetida).
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image as XlImage

if len(sys.argv) < 3:
    sys.exit("uso: python3 armar-excel.py <tablas.json> <salida.xlsx> [diagrama.png ...]")
spec = json.load(open(sys.argv[1], encoding="utf-8"))
salida = sys.argv[2]
imagenes = sys.argv[3:]

FUENTE = "Arial"
AZUL = "12287A"
GRIS = "F3F4F5"
f_base = Font(name=FUENTE, size=10)
f_neg = Font(name=FUENTE, size=10, bold=True)
f_tit = Font(name=FUENTE, size=16, bold=True, color=AZUL)
f_sub = Font(name=FUENTE, size=11, bold=True, color=AZUL)
f_cab = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
relleno_cab = PatternFill("solid", fgColor=AZUL)
relleno_gris = PatternFill("solid", fgColor=GRIS)
borde = Border(bottom=Side(style="thin", color="D0D3D7"))
envolver = Alignment(wrap_text=True, vertical="top")

# ── coherencia de la especificación ──
tablas = spec["tablas"]
nombres = [t["nombre"] for t in tablas]
rep = {n for n in nombres if nombres.count(n) > 1}
if rep:
    sys.exit(f"tablas repetidas: {sorted(rep)}")
cols = {t["nombre"]: {c["nombre"] for c in t["columnas"]} for t in tablas}
errores = []
def existe(ref):
    """«tabla.col» o «tabla.col1,col2» (clave compuesta): la tabla y cada columna existen."""
    tt, _, cc = ref.partition(".")
    return tt in cols and all(c.strip() in cols[tt] for c in cc.split(",") if c.strip())


for t in tablas:
    for c in t["columnas"]:
        # Una columna puede ser parte de varias claves foráneas: «a.x · b.y».
        for fk in (c.get("fk") or "").split(" · "):
            if fk.strip() and not existe(fk.strip()):
                errores.append(f"{t['nombre']}.{c['nombre']} → FK a {fk}, que no existe")
for r in spec.get("relaciones", []):
    for lado in ("desde", "hacia"):
        if not existe(r[lado]):
            errores.append(f"relación {r['desde']} → {r['hacia']}: {r[lado]} no existe")
    nd = len(r["desde"].partition(".")[2].split(",")); nh = len(r["hacia"].partition(".")[2].split(","))
    if nd != nh:
        errores.append(f"relación {r['desde']} → {r['hacia']}: {nd} columnas contra {nh}")
if errores:
    sys.exit("la especificación se contradice:\n  " + "\n  ".join(errores))
grupos = {g["id"]: g for g in spec["grupos"]}
sin_grupo = [t["nombre"] for t in tablas if t["grupo"] not in grupos]
if sin_grupo:
    sys.exit(f"tablas con un grupo que no está en «grupos»: {sin_grupo}")


def si_no(v):
    return "sí" if v else ""


def lista(v):
    return "\n".join(v) if isinstance(v, list) else (v or "")


def hoja_tabla(ws, nombre_tabla, cabeceras, filas, anchos, estilo="TableStyleMedium2"):
    """Escribe una tabla de Excel con cabecera, filtro, primera fila fija y anchos."""
    ws.append(cabeceras)
    for f in filas:
        ws.append(f)
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for celda in fila:
            celda.font = f_base
            celda.alignment = envolver
    for celda in ws[1]:
        celda.font = f_cab
    ref = f"A1:{get_column_letter(len(cabeceras))}{max(2, ws.max_row)}"
    tab = Table(displayName=nombre_tabla, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name=estilo, showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"


wb = Workbook()

# ── Léame ──
ws = wb.active
ws.title = "Léame"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 100
fila = 2
ws.cell(fila, 2, spec.get("titulo", "Modelo de datos · TRIM Hidrocarburos")).font = f_tit
fila += 1
ws.cell(fila, 2, spec.get("subtitulo", "Las tablas que guarda el backend y cómo se relacionan (PostgreSQL 16)")).font = Font(name=FUENTE, size=11, color="4D535A")
fila += 2
ws.cell(fila, 2, "CÓMO SE LEE").font = f_sub
fila += 1
for clave, texto in [
    ("Tablas", "Una fila por tabla: su grupo, para qué sirve, el tipo TypeScript del front que guarda, quién la escribe, qué endpoint la lee y sus reglas. «Columnas» y «Relaciones que salen» se cuentan con fórmulas sobre las otras hojas."),
    ("Columnas", "Una fila por columna: tipo PostgreSQL, si admite nulo, clave primaria (PK), clave foránea (FK: tabla.columna a la que apunta), único, restricción CHECK, descripción y un valor de ejemplo del demo. Use el filtro de la cabecera para ver una tabla."),
    ("Relaciones", "Cada clave foránea dicha en palabras: de qué columna a qué columna, su cardinalidad (N:1 = muchas filas de la primera apuntan a una de la segunda) y qué significa."),
    ("Filas de ejemplo", "Filas reales del demo (ids, fechas y valores de referencia/insumos), tabla por tabla. No son datos de un cliente."),
    ("Diagrama", "El diagrama entidad-relación de cada grupo, el mismo del capítulo 3 del documento."),
    ("Fuente de verdad", "El archivo que se ejecuta es esquema.sql (docs/backend/). Este libro se genera desde la misma especificación y se comprobó que coincide tabla por tabla y columna por columna."),
]:
    ws.cell(fila, 2, clave).font = f_neg
    c = ws.cell(fila, 3, texto)
    c.font = f_base
    c.alignment = envolver
    ws.row_dimensions[fila].height = 30
    fila += 1
fila += 1
ws.cell(fila, 2, "GRUPOS").font = f_sub
fila += 1
for g in spec["grupos"]:
    ws.cell(fila, 2, g["nombre"]).font = f_neg
    c = ws.cell(fila, 3, g["descripcion"])
    c.font = f_base
    c.alignment = envolver
    ws.row_dimensions[fila].height = 30
    fila += 1
fila += 1
ws.cell(fila, 2, "TOTALES").font = f_sub
fila += 1
ws.cell(fila, 2, "Tablas").font = f_neg
ws.cell(fila, 3, "=COUNTA(Tablas!B:B)-1").font = f_base
fila += 1
ws.cell(fila, 2, "Columnas").font = f_neg
ws.cell(fila, 3, "=COUNTA(Columnas!B:B)-1").font = f_base
fila += 1
ws.cell(fila, 2, "Relaciones").font = f_neg
ws.cell(fila, 3, "=COUNTA(Relaciones!A:A)-1").font = f_base
for r in range(fila - 2, fila + 1):
    ws.cell(r, 3).alignment = Alignment(horizontal="left")

# ── Tablas ──
ws = wb.create_sheet("Tablas")
filas = []
for i, t in enumerate(tablas, start=2):
    filas.append([
        grupos[t["grupo"]]["nombre"], t["nombre"], t["descripcion"], t.get("tipoTS") or "",
        t.get("quienEscribe") or "", lista(t.get("endpoints")), lista(t.get("reglas")),
        f'=COUNTIF(Columnas!$B:$B,B{i})', f'=COUNTIF(Relaciones!$A:$A,B{i})',
    ])
hoja_tabla(ws, "Tablas", ["Grupo", "Tabla", "Para qué sirve", "Tipo TS (front)", "Quién la escribe",
                          "Endpoints que la leen o escriben", "Reglas", "Columnas", "Relaciones que salen"],
           filas, [22, 26, 60, 26, 28, 40, 60, 11, 12])

# ── Columnas ──
ws = wb.create_sheet("Columnas")
filas = []
for t in tablas:
    for c in t["columnas"]:
        filas.append([grupos[t["grupo"]]["nombre"], t["nombre"], c["nombre"], c["tipo"],
                      "sí" if c.get("nulo") else "no", si_no(c.get("pk")), c.get("fk") or "",
                      (c["unico"] if isinstance(c.get("unico"), str) else si_no(c.get("unico"))),
                      c.get("check") or "", c.get("descripcion") or "",
                      "" if c.get("ejemplo") is None else str(c.get("ejemplo"))])
hoja_tabla(ws, "Columnas", ["Grupo", "Tabla", "Columna", "Tipo", "Nulo", "PK", "FK (tabla.columna)", "Único",
                            "CHECK", "Descripción", "Ejemplo"],
           filas, [20, 24, 24, 20, 7, 6, 28, 7, 34, 60, 30])

# ── Relaciones ──
ws = wb.create_sheet("Relaciones")
filas = []
for r in spec.get("relaciones", []):
    td, _, cd = r["desde"].partition(".")
    th, _, ch = r["hacia"].partition(".")
    filas.append([td, cd.replace(",", ", "), th, ch.replace(",", ", "), r.get("cardinalidad", "N:1"), r.get("descripcion", "")])
hoja_tabla(ws, "Relaciones", ["Tabla", "Columna", "Apunta a la tabla", "Columna", "Cardinalidad", "Qué significa"],
           filas, [26, 24, 26, 22, 12, 80])

# ── Filas de ejemplo ──
ws = wb.create_sheet("Filas de ejemplo")
ws.sheet_view.showGridLines = False
r = 1
ancho = {}
for t in tablas:
    filas_t = t.get("filas") or []
    if not filas_t:
        continue
    ws.cell(r, 1, t["nombre"]).font = f_sub
    r += 1
    claves = [c["nombre"] for c in t["columnas"] if any(c["nombre"] in f for f in filas_t)]
    for j, k in enumerate(claves, start=1):
        c = ws.cell(r, j, k)
        c.font = f_cab
        c.fill = relleno_cab
        ancho[j] = max(ancho.get(j, 10), min(40, len(k) + 2))
    r += 1
    for f in filas_t:
        for j, k in enumerate(claves, start=1):
            v = f.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False)
            c = ws.cell(r, j, "" if v is None else v)
            c.font = f_base
            c.alignment = Alignment(vertical="top")
            c.border = borde
            ancho[j] = max(ancho.get(j, 10), min(40, len(str(v)) + 2 if v is not None else 4))
        r += 1
    r += 1
for j, a in ancho.items():
    ws.column_dimensions[get_column_letter(j)].width = a

# ── Diagrama ──
if imagenes:
    ws = wb.create_sheet("Diagrama")
    ws.sheet_view.showGridLines = False
    fila = 1
    for ruta in imagenes:
        img = XlImage(ruta)
        escala = min(1.0, 1400 / img.width)
        img.width, img.height = int(img.width * escala), int(img.height * escala)
        ws.add_image(img, f"A{fila}")
        fila += int(img.height / 20) + 3

# Sin LibreOffice no hay valores calculados guardados: Excel, Numbers y Google Sheets
# recalculan todo al abrir con esta marca.
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(salida)

# Comprobación de las fórmulas sin LibreOffice: se simula cada COUNTIF sobre las hojas escritas
# y se compara con lo que la especificación dice. Una referencia a la columna equivocada falla aquí.
import re
from openpyxl import load_workbook
wb2 = load_workbook(salida)
col_b = [c.value for c in wb2["Columnas"]["B"][1:]]
rel_a = [c.value for c in wb2["Relaciones"]["A"][1:]]
malas = []
for fila in wb2["Tablas"].iter_rows(min_row=2):
    nombre, f_col, f_rel = fila[1].value, fila[7].value, fila[8].value
    esperado_col = next(len(t["columnas"]) for t in tablas if t["nombre"] == nombre)
    esperado_rel = sum(1 for r in spec.get("relaciones", []) if r["desde"].partition(".")[0] == nombre)
    m1 = re.fullmatch(r"=COUNTIF\(Columnas!\$B:\$B,B(\d+)\)", f_col)
    m2 = re.fullmatch(r"=COUNTIF\(Relaciones!\$A:\$A,B(\d+)\)", f_rel)
    if not m1 or not m2 or int(m1.group(1)) != fila[1].row or int(m2.group(1)) != fila[1].row:
        malas.append(f"{nombre}: fórmula mal armada")
        continue
    if col_b.count(nombre) != esperado_col or rel_a.count(nombre) != esperado_rel:
        malas.append(f"{nombre}: COUNTIF daría {col_b.count(nombre)}/{rel_a.count(nombre)}, se esperaba {esperado_col}/{esperado_rel}")
if malas:
    sys.exit("fórmulas que no cuadran:\n  " + "\n  ".join(malas))
print(f"{salida}: {len(tablas)} tablas, {sum(len(t['columnas']) for t in tablas)} columnas, "
      f"{len(spec.get('relaciones', []))} relaciones, {len(imagenes)} diagramas")
